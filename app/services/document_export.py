from __future__ import annotations

import difflib
from pathlib import Path

import fitz
from docx import Document
from pydantic import BaseModel, Field

from app.parsers.pdf_parser import page_text
from app.services.pdf_preview import apply_pdf_page_corrections


class ExportChunk(BaseModel):
    original_text: str
    revised_text: str
    page_index: int | None = Field(default=None, ge=0)
    start: int = Field(default=0, ge=0)


class ExportRequest(BaseModel):
    file_name: str = "document"
    chunks: list[ExportChunk]


STALE_REVIEW = "De controle hoort niet bij dit volledige document. Upload het document opnieuw."


def _revised_text(original: str, chunks: list[ExportChunk]) -> str:
    cursor = 0
    revised = []
    for chunk in chunks:
        end = cursor + len(chunk.original_text)
        if chunk.start != cursor or original[cursor:end] != chunk.original_text:
            raise ValueError(STALE_REVIEW)
        revised.append(chunk.revised_text)
        cursor = end
    if cursor != len(original):
        # Empty/whitespace-only files have no LanguageTool review chunks.
        if not chunks and not original.strip():
            return original
        raise ValueError(STALE_REVIEW)
    return "".join(revised)


def _export_docx(source: Path, chunks: list[ExportChunk], output: Path) -> None:
    document = Document(source)
    paragraphs = document.paragraphs
    original = "\n".join(paragraph.text for paragraph in paragraphs)
    revised = _revised_text(original, chunks)
    positions = []
    offset = 0
    for paragraph in paragraphs:
        positions.append((offset, offset + len(paragraph.text), paragraph))
        offset += len(paragraph.text) + 1

    edits = difflib.SequenceMatcher(None, original, revised, autojunk=False).get_opcodes()
    for tag, start, end, rstart, rend in reversed(edits):
        if tag == "equal":
            continue
        match = next(((pstart, paragraph) for pstart, pend, paragraph in positions
                      if pstart <= start <= end <= pend), None)
        if match is None or "\n" in revised[rstart:rend]:
            raise ValueError("Deze correctie wijzigt een alineagrens in het DOCX-document en kan niet met behoud van opmaak worden opgeslagen.")
        pstart, paragraph = match
        start -= pstart
        end -= pstart
        runs = paragraph.runs
        run_offset = 0
        affected = []
        for run in runs:
            run_end = run_offset + len(run.text)
            if (run_offset < end and run_end > start) or (start == end and run_offset <= start <= run_end):
                affected.append((run, run_offset))
                if start == end:
                    break
            run_offset = run_end
        if not affected:
            if not paragraph.text and start == end == 0:
                paragraph.add_run(revised[rstart:rend])
                continue
            raise ValueError("Deze DOCX-correctie kan niet met behoud van opmaak worden opgeslagen.")
        for index, (run, run_start) in enumerate(affected):
            left = max(0, start - run_start)
            right = min(len(run.text), end - run_start)
            run.text = run.text[:left] + (revised[rstart:rend] if index == 0 else "") + run.text[right:]
    document.save(output)


def export_document(source: Path, chunks: list[ExportChunk], output: Path) -> None:
    if source.suffix.lower() == ".pdf":
        with fitz.open(source) as document:
            by_page = {chunk.page_index: chunk for chunk in chunks}
            if len(chunks) != len(document) or set(by_page) != set(range(len(document))):
                raise ValueError(STALE_REVIEW)
            # Validate every page before making any changes.
            for index, page in enumerate(document):
                if by_page[index].original_text != page_text(page):
                    raise ValueError(STALE_REVIEW)
            for index, page in enumerate(document):
                chunk = by_page[index]
                if chunk.revised_text != chunk.original_text:
                    apply_pdf_page_corrections(page, chunk.revised_text)
            document.save(output, garbage=4, deflate=True)
    elif source.suffix.lower() == ".txt":
        original = source.read_text(encoding="utf-8")
        output.write_text(_revised_text(original, chunks), encoding="utf-8")
    elif source.suffix.lower() == ".docx":
        _export_docx(source, chunks, output)
    else:
        raise ValueError("Dit bestandstype kan niet worden geëxporteerd.")
