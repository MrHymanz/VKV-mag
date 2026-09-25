from pathlib import Path

from docx import Document

from app.models.document import ParsedDocument
from app.parsers.base import BaseParser


class DocxParser(BaseParser):
    supported_extensions = {".docx"}

    def parse(self, file_path: Path) -> ParsedDocument:
        document = Document(str(file_path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        text = "\n".join(paragraphs)
        return ParsedDocument(
            file_path=file_path,
            file_name=file_path.name,
            file_type="docx",
            text=text,
            metadata={"paragraph_count": len(document.paragraphs)},
            original_extension=file_path.suffix.lower(),
        )
