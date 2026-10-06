from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import main
from app.services.languagetool import LanguageToolClient
from app.services.spacing_check import add_docx_spacing_issues


def test_spacing_skips_indentation_tabs_and_newlines():
    text = '  Kop\nTekst  met   spaties.\n\t  Begin\nRegel  \nVolgende\t\twoord'
    chunks = [{'id': 1, 'start': 0, 'text': text, 'issues': []}]
    issues = add_docx_spacing_issues(text, chunks)[0]['issues']
    assert [text[i['offset']:i['offset'] + i['length']] for i in issues] == ['  ', '   ']
    assert all(not i['accepted'] for i in issues)


def test_space_run_at_chunk_boundary_and_duplicate_issue():
    text = 'Dit   klopt.'
    duplicate = {'offset': 3, 'length': 3, 'replacements': [{'value': ' '}], 'rule_id': 'WHITESPACE_RULE'}
    chunks = [{'id': 1, 'start': 0, 'text': text[:4], 'issues': [duplicate]},
              {'id': 2, 'start': 4, 'text': text[4:], 'issues': []}]
    result = add_docx_spacing_issues(text, chunks)
    assert len(result) == 1
    assert result[0]['text'] == text
    assert len(result[0]['issues']) == 1
    assert result[0]['issues'][0]['rule_id'] == 'VKV_DOUBLE_SPACES'


@pytest.mark.parametrize('sensitivity', ['cautious', 'normal', 'strict'])
def test_docx_upload_and_export_spacing_preserves_runs(tmp_path, monkeypatch, sensitivity):
    monkeypatch.setattr(main, 'UPLOAD_DIR', tmp_path)
    monkeypatch.setattr(main, 'TEMP_DIR', tmp_path)
    monkeypatch.setattr(LanguageToolClient, '_request_chunk', lambda *a, **kw: [])
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Dit ').bold = True
    paragraph.add_run(' is mooi.').italic = True
    data = BytesIO()
    document.save(data)
    client = TestClient(main.app)
    response = client.post('/upload', files={'file': ('test.docx', data.getvalue())}, data={'sensitivity': sensitivity})
    assert response.status_code == 200
    review = response.json()
    issue = review['issues'][0]
    assert issue['rule_id'] == 'VKV_DOUBLE_SPACES'
    assert issue['accepted'] is False
    original = review['chunks'][0]['text']
    revised = original[:issue['offset']] + ' ' + original[issue['offset'] + issue['length']:]
    response = client.post('/download/' + review['stored_name'], json={
        'file_name': 'test.docx', 'chunks': [{'start': 0, 'original_text': original, 'revised_text': revised}]})
    assert response.status_code == 200
    output = Document(BytesIO(response.content))
    assert output.paragraphs[0].text == 'Dit is mooi.'
    assert output.paragraphs[0].runs[0].bold
    assert output.paragraphs[0].runs[1].italic
    txt = client.post('/upload', files={'file': ('test.txt', b'Dit  is mooi.')})
    assert txt.json()['issues'] == []
