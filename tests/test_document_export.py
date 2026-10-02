from io import BytesIO
from urllib.parse import unquote

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import main
from app.parsers.pdf_parser import page_text


@pytest.fixture
def client(tmp_path, monkeypatch):
    uploads = tmp_path / 'uploads'
    exports = tmp_path / 'exports'
    uploads.mkdir()
    exports.mkdir()
    monkeypatch.setattr(main, 'UPLOAD_DIR', uploads)
    monkeypatch.setattr(main, 'TEMP_DIR', exports)
    return TestClient(main.app)


def test_download_pdf_uses_all_pages_and_preserves_rejected_edits(client):
    source = main.UPLOAD_DIR / 'magazine.pdf'
    with fitz.open() as document:
        for text in ['De redacteur bedoeld dat.', 'De mijne is zeker betrouwbaar.', 'Een spelfoutt.']:
            page = document.new_page(width=400, height=200)
            page.insert_text((40, 70), text, fontsize=12)
            page.draw_rect(fitz.Rect(300, 100, 340, 140), fill=(1, 0, 0))
        document.new_page(width=400, height=200)  # Keep blank pages too.
        document.set_metadata({'title': 'VKV test'})
        document.set_toc([[1, 'Artikel', 1], [1, 'Vervolg', 3]])
        document.save(source)
        chunks = [{'page_index': index, 'original_text': page_text(page),
                   'revised_text': page_text(page)} for index, page in enumerate(document)]
    original_bytes = source.read_bytes()
    chunks[0]['revised_text'] = 'De redacteur bedoelt dat.'
    chunks[2]['revised_text'] = 'Een spelfout.'
    response = client.post('/download/magazine.pdf', json={'file_name': 'VKV oktober.pdf', 'chunks': chunks})
    assert response.status_code == 200, response.text
    assert 'application/pdf' in response.headers['content-type']
    assert 'VKV oktober-gecorrigeerd.pdf' in unquote(response.headers['content-disposition'])
    with fitz.open(stream=response.content, filetype='pdf') as result:
        assert len(result) == 4
        assert 'bedoelt' in result[0].get_text() and 'bedoeld' not in result[0].get_text()
        assert 'De mijne' in result[1].get_text()
        assert 'spelfout.' in result[2].get_text() and 'spelfoutt' not in result[2].get_text()
        assert result[3].get_text() == ''
        assert result.metadata['title'] == 'VKV test'
        assert result.get_toc() == [[1, 'Artikel', 1], [1, 'Vervolg', 3]]
        assert result[0].rect.width == 400  # Single pages, not a comparison spread.
        drawings = result[0].get_drawings()
        assert len(drawings) == 1  # Artwork remains; no correction highlights.
        assert drawings[0]['fill'] == (1, 0, 0)
    assert source.read_bytes() == original_bytes

    # Download again after unchecking: do not serve the previous export.
    chunks[0]['revised_text'] = chunks[0]['original_text']
    response = client.post('/download/magazine.pdf', json={'chunks': chunks})
    with fitz.open(stream=response.content, filetype='pdf') as result:
        assert 'bedoeld' in result[0].get_text()
        assert 'spelfout.' in result[2].get_text()


def test_download_pdf_keeps_word_breaks(client):
    source = main.UPLOAD_DIR / 'split.pdf'
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((40, 50), 'Betrouwba-')
        page.insert_text((40, 70), 're auto met spelfoutt.')
        original = page_text(page)
        document.save(source)
    response = client.post('/download/split.pdf', json={'chunks': [{
        'page_index': 0, 'original_text': original,
        'revised_text': 'Betrouwbare auto met spelfout.',
    }]})
    assert response.status_code == 200
    with fitz.open(stream=response.content, filetype='pdf') as result:
        text = result[0].get_text()
        assert 'Betrouwba-' in text and 're auto met' in text
        assert 'spelfoutt' not in text


@pytest.mark.parametrize('chunks', [[], [{'page_index': 0, 'original_text': 'wrong', 'revised_text': 'new'}]])
def test_incomplete_or_stale_pdf_review_rejected(client, chunks):
    with fitz.open() as document:
        document.new_page().insert_text((40, 50), 'Origineel')
        document.save(main.UPLOAD_DIR / 'original.pdf')
    response = client.post('/download/original.pdf', json={'chunks': chunks})
    assert response.status_code == 409
    assert not list(main.TEMP_DIR.iterdir())


def test_old_download_route_does_not_mislabel_original(client):
    response = client.get('/download/original.pdf')
    assert response.status_code == 409
    assert 'Ververs' in response.json()['detail']


def test_txt_export_keeps_chunk_whitespace_and_latest_choices(client):
    source = main.UPLOAD_DIR / 'text.txt'
    source.write_text('De mijne.\nEen spelfoutt.\n', encoding='utf-8')
    response = client.post('/download/text.txt', json={'chunks': [
        {'start': 0, 'original_text': 'De mijne.\n', 'revised_text': 'De mijne.\n'},
        {'start': 10, 'original_text': 'Een spelfoutt.\n', 'revised_text': 'Een spelfout.\n'},
    ]})
    assert response.status_code == 200
    assert response.text == 'De mijne.\nEen spelfout.\n'
    assert source.read_text(encoding='utf-8') == 'De mijne.\nEen spelfoutt.\n'


def test_docx_export_preserves_runs_and_untouched_paragraphs(client):
    source = main.UPLOAD_DIR / 'text.docx'
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Een ').bold = True
    paragraph.add_run('spel').italic = True
    paragraph.add_run('foutt.').italic = True
    document.add_paragraph('De mijne blijft staan.')
    document.save(source)
    response = client.post('/download/text.docx', json={'chunks': [{
        'original_text': 'Een spelfoutt.\nDe mijne blijft staan.',
        'revised_text': 'Een spelfout.\nDe mijne blijft staan.',
    }]})
    assert response.status_code == 200, response.text
    result = Document(BytesIO(response.content))
    assert result.paragraphs[0].text == 'Een spelfout.'
    assert result.paragraphs[0].runs[0].bold is True
    assert result.paragraphs[0].runs[1].italic is True
    assert result.paragraphs[1].text == 'De mijne blijft staan.'
